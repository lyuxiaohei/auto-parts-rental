# -*- coding: utf-8 -*-
# 修复甘特条形 CF 公式列引用（$G/$F、$E/$D）+ 标题行数
import re
from copy import deepcopy
from openpyxl import load_workbook
from openpyxl.formatting.formatting import ConditionalFormattingList

F = 'P1-R03-工期评估表-20260915.xlsx'
wb = load_workbook(F)
fail = []
for nm, axis_col in [('甘特图-任务级', 'H'), ('甘特图-功能点级', 'F')]:
    ws = wb[nm]
    captured = [r for rng in ws.conditional_formatting for r in rng.rules]
    bars = [r for r in captured if r.formula and 'AND(' in str(r.formula[0])]
    ms = [r for r in captured if r.formula and 'OR(' in str(r.formula[0])][0]
    td = [r for r in captured if r.formula and 'TODAY' in str(r.formula[0])][0]
    bar_sqs = []
    for rng in ws.conditional_formatting:
        for r in rng.rules:
            if r.formula and 'AND(' in str(r.formula[0]):
                bar_sqs.append(str(rng.sqref))
    if not bar_sqs:
        fail.append(nm + ' 无条形规则'); continue
    ref = bars[0]
    if not (isinstance(ref.formula, list)):
        ref.formula = [str(ref.formula)]
    cols = re.findall(r'\$([A-Z]+)\d+', str(ref.formula[0]))
    if len(cols) < 3:
        fail.append('%s 列提取 %s from %s' % (nm, cols, ref.formula[0])); continue
    endc, startc = cols[1], cols[2]
    full_last = bar_sqs[0].split(':')[1].lstrip('HF')
    full_last = re.sub(r'^[A-Z]+', '', bar_sqs[0].split(':')[1])
    full_range = '%s5:%s' % (axis_col, bar_sqs[0].split(':')[1].split('5:')[-1] if '5:' in bar_sqs[0] else bar_sqs[0].split(':')[1])
    ws.conditional_formatting = ConditionalFormattingList()
    ws.conditional_formatting.add(full_range, ms)
    ws.conditional_formatting.add(full_range, td)
    for sq in bar_sqs:
        anchor = int(re.findall(r'[A-Z]+(\d+)$', sq.split(':')[0])[0])
        r2 = deepcopy(ref)
        r2.formula = ['AND(%s$4<>"",%s$4<=$%s%d,%s$4>=$%s%d)' % (axis_col, axis_col, endc, anchor, axis_col, startc, anchor)]
        ws.conditional_formatting.add(sq, r2)
wb['甘特图-任务级']['A1'] = '甘特图 · 任务级（67 行全量·0915 增补）'
wb.save(F)

wb2 = load_workbook(F)
ok = True
for nm in ['甘特图-任务级', '甘特图-功能点级']:
    ws = wb2[nm]
    n = 0
    for rng in ws.conditional_formatting:
        for r in rng.rules:
            n += 1
            if r.formula and 'AND(' in str(r.formula[0]):
                print('%s %s → %s' % (nm, str(rng.sqref), r.formula[0]))
    print('%s 规则总数=%d' % (nm, n))
    ok = ok and n == 7
print('失败 %d 条｜七规则齐=%s' % (len(fail), ok))
for x in fail: print('  ' + x)
