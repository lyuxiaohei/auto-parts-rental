# -*- coding: utf-8 -*-
"""Dump full structure of the schedule xlsx for replication in Tencent Docs."""
import json
import sys
from datetime import datetime, date
from openpyxl import load_workbook

PATH = r"D:\工作台-吕道远\5-【ACTIVE】汽车物流包装租赁\P1-R03-工期评估表-20260915.xlsx"

wb = load_workbook(PATH, data_only=True)

def fmt(v):
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d") if not isinstance(v, datetime) or v.hour == 0 and v.minute == 0 else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, float) and v == int(v):
        # keep display-like value
        return int(v)
    return str(v) if not isinstance(v, (int, float)) else v

result = {"sheets": []}
for ws in wb.worksheets:
    merged = [str(r) for r in ws.merged_cells.ranges]
    # determine real used range
    max_row = ws.max_row
    max_col = ws.max_column
    rows = []
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        rows.append([fmt(c.value) for c in row])
    # trim trailing empty rows
    while rows and all(v is None or (isinstance(v, str) and v.strip() == "") for v in rows[-1]):
        rows.pop()
        max_row -= 1
    # trim trailing empty cols
    if rows:
        width = max(len(r) for r in rows)
        while width > 0 and all(len(r) < width or r[width-1] is None or (isinstance(r[width-1], str) and r[width-1].strip() == "") for r in rows):
            width -= 1
        rows = [r[:width] for r in rows]
        max_col = width
    # bold info for header row(s): detect bold on first row
    bold_first_row = []
    if rows:
        for c in ws[1]:
            f = c.font
            bold_first_row.append(bool(f and f.bold))
    result["sheets"].append({
        "name": ws.title,
        "max_row": len(rows),
        "max_col": max_col,
        "merged": merged,
        "freeze": ws.freeze_panes,
        "bold_first_row": bold_first_row,
        "rows": rows,
    })

print(json.dumps(result, ensure_ascii=False, indent=1))
