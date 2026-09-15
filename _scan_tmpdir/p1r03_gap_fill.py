# -*- coding: utf-8 -*-
# 工期表补漏：D-120 企微/飞书组织架构同步（B0·V1.1·人日待评估）插行 + FP3-03 备注 D-144
import datetime
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

F = 'P1-R03-工期评估表-20260915.xlsx'
wb = load_workbook(F)
fp = wb['功能点评估（全栈）']; hs = wb['汇总']; t = wb['甘特图-任务级']
fail = []
def chk(c, m):
    if not c: fail.append(m)

# ---------- 1. 功能点表：B0 尾（r15）插入 D-120 ----------
chk(fp.cell(row=14, column=4).value.startswith('移动端 H5'), 'r14=%r' % fp.cell(row=14, column=4).value)
styles = [copy(fp.cell(row=14, column=c)._style) for c in range(1, 14)]
fp.insert_rows(15, 1)
for c, st in enumerate(styles, 1): fp.cell(row=15, column=c)._style = st
vals = {1: 'B0', 2: '公共底座与移动端', 3: '（待补）',
        4: '企微/飞书组织架构同步：部门与人员同步＋账号映射（D-120 刚性需求·接口成熟）',
        5: '全栈', 6: '后台通道（无独立页）', 7: 'V1.1', 8: None,
        9: '=H15*(1+汇总!$C$46)', 10: '吕道远', 11: None, 12: None,
        13: 'D-120（第4次沟通 1:06:39 王琳总·拍板需工期评估）；人日待开发评估'}
for c, v in vals.items(): fp.cell(row=15, column=c).value = v
fp.cell(row=15, column=1).fill = PatternFill('solid', start_color='FFDEEBF7')
fp.cell(row=15, column=7).fill = PatternFill('solid', start_color='FFE2EFDA')

# ---------- 2. FP3-03（租赁出库单）备注补 D-144 ----------
row33 = None
for r in range(4, 72):
    if str(fp.cell(row=r, column=4).value or '').startswith('租赁出库单'):
        row33 = r; break
chk(row33, 'FP3-03 未找到')
m = fp.cell(row=row33, column=13).value or ''
if 'D-144' not in m:
    fp.cell(row=row33, column=13).value = (m + '；含出货单打印页（D-144·打印模板+取数）').strip('；')

# ---------- 3. 汇总 D/E 块区间 +1（B1 起全部下移一行） ----------
BLOCKS = [('A1',4,7),('A2',8,9),('B0',10,15),('B1',16,20),('B2',21,25),('B3',26,37),('B4',38,42),
          ('B5',43,47),('B6',48,55),('B7',56,58),('B8',59,63),('C1',64,65),('C2',66,68),('C3',69,71)]
S = "'功能点评估（全栈）'!"
for (w, a, b), r in zip(BLOCKS, range(4, 18)):
    chk(hs.cell(row=r, column=1).value == w, '汇总r%d=%s' % (r, w))
    hs.cell(row=r, column=4).value = '=IF(COUNT(%sK%d:K%d)=0,"",MIN(%sK%d:K%d))' % (S, a, b, S, a, b)
    hs.cell(row=r, column=5).value = '=IF(COUNT(%sL%d:L%d)=0,"",MAX(%sL%d:L%d))' % (S, a, b, S, a, b)
for r in range(4, 18):
    f_ = hs.cell(row=r, column=3).value
    if isinstance(f_, str) and 'SUMIF' in f_:
        hs.cell(row=r, column=3).value = f_.replace('$A$4:$A$75', '$A$4:$A$76').replace('$H$4:$H$75', '$H$4:$H$76')

# ---------- 4. 甘特任务级：V1.1 区首行（移动端）后插 D-120 行 + CF 区间顺延 ----------
mov = None
for r in range(6, 80):
    if str(t.cell(row=r, column=3).value or '').startswith('移动端 H5'):
        mov = r; break
chk(mov, '甘特移动端行未找到')
gstyles = [copy(t.cell(row=mov, column=c)._style) for c in range(1, 86)]
t.insert_rows(mov + 1, 1)
for c, st in enumerate(gstyles, 1): t.cell(row=mov + 1, column=c)._style = st
for c, v in enumerate(['B0', '（待补）', '企微/飞书组织架构同步（部门/人员/账号映射·D-120）', 'V1.1', None, None, None], 1):
    t.cell(row=mov + 1, column=c).value = v
# 带行重定位 + 重合并 + CF 区间重建
def find_band(prefix):
    for r in range(5, 95):
        if str(t.cell(row=r, column=1).value or '').startswith(prefix): return r
bands = {k: find_band(k) for k in ['前置（', 'V1.0 买卖版', 'V1.1 租赁版', 'V2.0 深化版', '各版支撑']}
chk(all(bands.values()), '带定位 %s' % bands)
for k in ['V2.0 深化版', '各版支撑']:
    r = bands[k]
    if 'A%d:G%d' % (r, r) not in [str(m) for m in t.merged_cells.ranges]:
        t.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
last_data = bands['各版支撑'] + 8
from openpyxl.formatting.formatting import ConditionalFormattingList
from copy import deepcopy
captured = [r2 for rng in t.conditional_formatting for r2 in rng.rules]
ms = [r2 for r2 in captured if r2.formula and 'OR(' in str(r2.formula[0])][0]
td = [r2 for r2 in captured if r2.formula and 'TODAY' in str(r2.formula[0])][0]
ref = [r2 for r2 in captured if r2.formula and 'AND(' in str(r2.formula[0])][0]
spans = [(bands['前置（'] + 1, bands['V1.0 买卖版'] - 1, 'FFD9E2F3'),
         (bands['V1.0 买卖版'] + 1, bands['V1.1 租赁版'] - 1, 'FFDDEBF7'),
         (bands['V1.1 租赁版'] + 1, bands['V2.0 深化版'] - 1, 'FFE2EFDA'),
         (bands['V2.0 深化版'] + 1, bands['各版支撑'] - 1, 'FFFCE4D6'),
         (bands['各版支撑'] + 1, last_data, 'FFE7E6E6')]
t.conditional_formatting = ConditionalFormattingList()
t.conditional_formatting.add('H5:CG%d' % last_data, ms)
t.conditional_formatting.add('H5:CG%d' % last_data, td)
for (a, b, _) in spans:
    r2 = deepcopy(ref)
    r2.formula = ['AND(H$4<>"",H$4<=$G%d,H$4>=$F%d)' % (a, a)]
    t.conditional_formatting.add('H%d:CG%d' % (a, b), r2)
t['A1'] = '甘特图 · 任务级（%d 行全量·0915 增补）' % (last_data - 10)

tot = sum(fp.cell(row=r, column=8).value for r in range(4, 72) if isinstance(fp.cell(row=r, column=8).value, (int, float)))
chk(abs(tot - 153.75) < 1e-9, '合计 %.2f' % tot)
wb.save(F)
print('失败 %d 条｜合计=%.2f｜甘特带=%s｜末数据行=%d' % (len(fail), tot, bands, last_data))
for x in fail: print('  ' + x)
