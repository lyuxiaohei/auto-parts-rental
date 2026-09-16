# -*- coding: utf-8 -*-
"""G41 T8：P1-R03 xlsx 注记——zip 级精准注入（不动 drawing/media/其余 sheet）。"""
import zipfile, shutil, re, os

ROOT = r"D:\工作台-吕道远\5-【ACTIVE】汽车物流包装租赁"
SRC = ROOT + r"\P1-R03-工期评估表-20260915.xlsx"
TMP = ROOT + r"\_scan_tmpdir\_p1r03_new.xlsx"

z = zipfile.ZipFile(SRC)
wbxml = z.read("xl/workbook.xml").decode("utf-8")
rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")

# 1. 定位「功能点评估（全栈）」sheet 的 XML 路径
m = re.search(r'<sheet[^>]*name="功能点评估（全栈）"[^>]*r:id="(rId\d+)"', wbxml)
assert m, "sheet 名未找到"
rid = m.group(1)
m2 = re.search(r'<Relationship[^>]*Id="%s"[^>]*Target="([^"]+)"' % rid, rels)
assert m2, "rels 未找到"
sheet_path = "xl/" + m2.group(1).lstrip("/")
print("目标 sheet XML：", sheet_path)

sx = z.read(sheet_path).decode("utf-8")
# 2. 实际末行号
rowids = [int(r) for r in re.findall(r'<row r="(\d+)"', sx)]
last = max(rowids)
print("当前末行：", last)
note_r = last + 1
note = ("<row r=\"%d\">" % note_r
        + "<c r=\"A%d\" t=\"inlineStr\"><is><t>【G41 注记·2026-09-16】</t></is></c>" % note_r
        + "<c r=\"B%d\" t=\"inlineStr\"><is><t>G37 转移出库／G38 字段口径／G39 按天计租 三任务增量工作量按 D-135 只注记不重算；本表数字维持 V20260915（D-94）口径。增量明细见 P2-R01 决策台账 D-146~D-150 与对应任务档（agent-handoff/）。</t></is></c>" % note_r
        + "</row>")
assert "<sheetData/>" not in sx
assert sx.count("</sheetData>") == 1
sx_new = sx.replace("</sheetData>", note + "</sheetData>")
# 3. dimension 同步（容忍缺失）
dm = re.search(r'<dimension ref="([A-Z]+\d+:[A-Z]+\d+)"/?>', sx_new)
if dm:
    cur_end = dm.group(1).split(":")[1]
    col = re.match(r"([A-Z]+)(\d+)", cur_end).group(1)
    sx_new = sx_new.replace(dm.group(0), '<dimension ref="A1:%s%d"/>' % (col, note_r))
    print("dimension 已更新 → A1:%s%d" % (col, note_r))

# 4. 重写 zip（其余条目字节原样）
zo = zipfile.ZipFile(TMP, "w", zipfile.ZIP_DEFLATED)
for item in z.infolist():
    data = z.read(item.filename)
    if item.filename == sheet_path:
        data = sx_new.encode("utf-8")
    zo.writestr(item, data)
zo.close()
z.close()

# 5. 校验：媒体字节不变＋新行可读
za = zipfile.ZipFile(SRC); zb = zipfile.ZipFile(TMP)
same = all(za.read(n) == zb.read(n) for n in za.namelist() if n != sheet_path)
print("其余条目字节级一致：", same)
import openpyxl
wb = openpyxl.load_workbook(TMP)
ws = wb["功能点评估（全栈）"]
print("读回注记：", str(ws.cell(note_r, 1).value)[:20], "|", str(ws.cell(note_r, 2).value)[:40])
print("图片仍挂载：", len(ws._images) if hasattr(ws, "_images") else "n/a")
wsl = wb["版本交付阶梯图"]
print("阶梯图 sheet 图片数：", len(wsl._images))
assert same and str(ws.cell(note_r, 1).value).startswith("【G41 注记")
za.close(); zb.close(); wb.close()

# 6. 原子替换＋清理探针
os.replace(TMP, SRC)
os.remove(ROOT + r"\_scan_tmpdir\_p1r03_probe.xlsx")
print("已原子替换 P1-R03-工期评估表-20260915.xlsx")
