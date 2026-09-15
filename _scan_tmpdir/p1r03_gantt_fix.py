# -*- coding: utf-8 -*-
# 修正国庆落空的排期（销售退货 10-01~02→09-23~24；V1.0 开发窗收尾 10-01→09-30）并重画两张甘特图
import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

F = 'P1-R03-工期评估表-20260915.xlsx'
D = lambda m, d: datetime.datetime(2026, m, d)
wb = load_workbook(F)
fp = wb['功能点评估（全栈）']; hs = wb['汇总']
fail = []
def chk(c, m):
    if not c: fail.append(m)

chk(fp.cell(row=46, column=4).value.startswith('销售退货单'), '行46 非销售退货')
fp.cell(row=46, column=10).value = D(9, 23)
fp.cell(row=46, column=11).value = D(9, 24)
chk(hs['G23'].value == D(10, 1), 'G23=%r' % hs['G23'].value)
hs['G23'] = D(9, 30)

# ---- 重画甘特（与全量脚本同一逻辑） ----
def workdays(d1, d2, skip=frozenset([datetime.date(2026, 10, d) for d in range(1, 8)])):
    out, d = [], d1
    while d <= d2:
        if d.weekday() < 5 and d not in skip: out.append(d)
        d += datetime.timedelta(days=1)
    return out

AXIS = workdays(datetime.date(2026, 9, 2), datetime.date(2026, 12, 25))
COLOR = {'前置': 'FFB9C0CC', 'V1.0': 'FF6E9BD1', 'V1.1': 'FF3E78C2', 'V2.0': 'FF1B5FAD', '各版': 'FFFFC000'}
F_TITLE = Font(bold=True, size=13); F_NOTE = Font(size=9, color='FF6B7280')
F_HDR = Font(bold=True, size=9); F_DAT = Font(size=9); F_DAY = Font(size=8)
FILL_HDR = PatternFill('solid', start_color='FFEDF2F9')
AL_C = Alignment(horizontal='center')

rows = []
for r in range(4, 71):
    if not fp.cell(row=r, column=1).value: continue
    rows.append(dict(wbs=fp.cell(row=r, column=1).value, fpid=fp.cell(row=r, column=3).value,
                     task=fp.cell(row=r, column=4).value, ver=fp.cell(row=r, column=7).value,
                     pd=fp.cell(row=r, column=8).value, st=fp.cell(row=r, column=10).value,
                     en=fp.cell(row=r, column=11).value))
chk(len(rows) == 67, '行数 %d' % len(rows))

def build_sheet(name, idx, meta, data, note):
    if name in wb.sheetnames: wb.remove(wb[name])
    ws = wb.create_sheet(name, idx)
    n0 = len(meta)
    ws.cell(row=1, column=1).value = ('甘特图 · 任务级（%d 行全量）' if '任务' in name else '甘特图 · WBS × 版本（%d 行概览版）') % len(data)
    ws.cell(row=1, column=1).font = F_TITLE
    ws.cell(row=2, column=1).value = note; ws.cell(row=2, column=1).font = F_NOTE
    for c, h in enumerate(meta, 1):
        cell = ws.cell(row=3, column=c, value=h); cell.font = F_HDR; cell.fill = FILL_HDR; cell.alignment = AL_C
    prev = None
    for i, d in enumerate(AXIS):
        c = n0 + 1 + i
        cell = ws.cell(row=4, column=c, value=d.day); cell.font = F_DAY; cell.alignment = AL_C; cell.fill = FILL_HDR
        if prev is None or d.month != prev.month:
            mc = ws.cell(row=3, column=c, value='%d月' % d.month); mc.font = F_HDR; mc.fill = FILL_HDR; mc.alignment = AL_C
        prev = d
    r = 5
    band = ws.cell(row=r, column=1, value='前置（09-02 ~ 09-18 · 需求+原型+底座）'); band.font = F_HDR; band.fill = FILL_HDR
    r += 1
    for row in data:
        vals = [row['wbs'], row.get('fpid', ''), row['task'], row['ver'], row['pd'], row['st'], row['en']][:n0]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v); cell.font = F_DAT
            if isinstance(v, datetime.datetime): cell.number_format = 'MM-DD'
            if c >= 4 or c == 1: cell.alignment = AL_C
        color = COLOR.get(row['ver'])
        if color and isinstance(row['st'], datetime.datetime):
            s = row['st'].date()
            e = row['en'].date() if isinstance(row['en'], datetime.datetime) else s
            bar = PatternFill('solid', start_color=color)
            for i, d in enumerate(AXIS):
                if s <= d <= e: ws.cell(row=r, column=n0 + 1 + i).fill = bar
        r += 1
    for c in range(1, n0 + 1):
        ws.column_dimensions[get_column_letter(c)].width = {1: 5, 2: 9, 3: 44, 4: 7, 5: 6, 6: 8, 7: 8}.get(c, 9)
    for i in range(len(AXIS)):
        ws.column_dimensions[get_column_letter(n0 + 1 + i)].width = 2.64
    ws.freeze_panes = ws.cell(row=5, column=n0 + 1)

note = '口径：横轴 %d 个工作日（09-02~12-25·周末与国庆 10-01~07 不排）；09-15 增补口径（+12.5 人日·153.75），节点 10-20/11-24/12-14，人日待开发重估后刷新' % len(AXIS)
build_sheet('甘特图-任务级', 3, ['WBS', '功能点', '任务', '版本', '人日', '开始', '结束'], rows, note)

agg_pre = [dict(wbs='前置', fpid='', task='需求调研与 PRD（四轮走查+计费定案+评审冻结·09-18）', ver='前置', pd=6, st=D(9, 2), en=D(9, 18)),
           dict(wbs='前置', fpid='', task='原型改造与定稿（G31~G39·128+ HTML）', ver='前置', pd=1, st=D(9, 14), en=D(9, 18))]
b0 = [x for x in rows if x['wbs'] == 'B0']
agg_pre.append(dict(wbs='前置', fpid='', task='公共底座（脚手架/登录/通用组件）', ver='前置',
                    pd=sum(x['pd'] for x in b0 if x['ver'] == '前置'),
                    st=min(x['st'] for x in b0 if x['ver'] == '前置'),
                    en=max(x['en'] for x in b0 if x['ver'] == '前置')))
groups = {}
for x in [y for y in rows if y['wbs'] not in ('A1', 'A2', 'B0')]:
    k = (x['wbs'], x['ver'])
    g = groups.setdefault(k, dict(wbs=x['wbs'], fpid='', task=x['task'], ver=x['ver'], pd=0, st=None, en=None))
    g['pd'] += x['pd'] or 0
    if isinstance(x['st'], datetime.datetime): g['st'] = x['st'] if g['st'] is None else min(g['st'], x['st'])
    if isinstance(x['en'], datetime.datetime): g['en'] = x['en'] if g['en'] is None else max(g['en'], x['en'])
order = {w: i for i, w in enumerate(['B1', 'B2', 'B3', 'B4', 'B5', 'B6', 'B7', 'B8', 'C1', 'C2', 'C3'])}
vord = {'V1.0': 0, 'V1.1': 1, 'V2.0': 2, '各版': 3}
gl = [groups[k] for k in sorted(groups, key=lambda k: (order.get(k[0], 99), vord.get(k[1], 9)))]
build_sheet('甘特图-功能点级', 2, ['WBS', '模块 / 内容', '人日', '开始', '结束'], agg_pre + gl, note)

wb.save(F)
print('失败 %d 条｜轴 %d 天' % (len(fail), len(AXIS)))
for x in fail: print('  ' + x)
