# -*- coding: utf-8 -*-
# P1-R03 全量更新 → 20260915 版：甘特图两张重排 + V2.0/C 行日期顺延 + 汇总残留修正 + 归档 0914 + 更名 0915
import shutil, datetime, io
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

SRC = 'P1-R03-工期评估表-20260914.xlsx'
DST = 'P1-R03-工期评估表-20260915.xlsx'
wb = load_workbook(SRC)
fp = wb['功能点评估（全栈）']; hs = wb['汇总']
fail = []
def chk(c, m):
    if not c: fail.append(m)

D = lambda m, d: datetime.datetime(2026, m, d)

# ---------- 1. V2.0 功能点行日期 +7 天（≈+5 工作日·窗口 11-13~11-30） ----------
V2_SHIFT = []
for r in range(4, 71):
    if fp.cell(row=r, column=7).value == 'V2.0':
        j = fp.cell(row=r, column=10).value; k = fp.cell(row=r, column=11).value
        if isinstance(j, datetime.datetime):
            fp.cell(row=r, column=10).value = j + datetime.timedelta(days=7)
            fp.cell(row=r, column=11).value = (k + datetime.timedelta(days=7)) if isinstance(k, datetime.datetime) else k
            V2_SHIFT.append(r)
chk(len(V2_SHIFT) == 4, 'V2.0 行数 %d≠4' % len(V2_SHIFT))

# ---------- 2. C1-C3 测试/迁移/部署行对齐新轮次（10-13 / 11-16 / 12-08 启动） ----------
CDATES = {  # r: (开始, 结束)
 63: (D(10,13), D(12,11)), 64: (D(10,19), D(12,14)),
 65: (D(10,13), D(12,9)),  66: (D(10,16), D(12,10)), 67: (D(10,19), D(12,11)),
 68: (D(10,19), D(12,11)), 69: (D(10,20), D(12,14)), 70: (D(10,20), D(12,25)),
}
for r, (a, b) in CDATES.items():
    chk(fp.cell(row=r, column=1).value in ('C1','C2','C3'), '行%d 非 C 行'%r)
    fp.cell(row=r, column=10).value = a
    fp.cell(row=r, column=11).value = b

# ---------- 3. 汇总残留（无AI 212→约219） ----------
h21 = hs['H21'].value; chk('212' in str(h21), 'H21=%r' % h21)
hs['H21'] = str(h21).replace('212人日', '约219人日')
r4 = hs['C41'].value; chk('212' in str(r4), 'R4=%r' % r4)
hs['C41'] = str(r4).replace('212 人日', '约 219 人日')

# ---------- 4. 甘特图重排 ----------
def workdays(d1, d2, skip=frozenset([datetime.date(2026,10,d) for d in range(1,8)])):
    out, d = [], d1
    while d <= d2:
        if d.weekday() < 5 and d not in skip: out.append(d)
        d += datetime.timedelta(days=1)
    return out

AXIS = workdays(datetime.date(2026,9,2), datetime.date(2026,12,25))
COLOR = {'前置':'FFB9C0CC','V1.0':'FF6E9BD1','V1.1':'FF3E78C2','V2.0':'FF1B5FAD','各版':'FFFFC000'}
F_TITLE = Font(bold=True, size=13); F_NOTE = Font(size=9, color='FF6B7280')
F_HDR = Font(bold=True, size=9); F_DAT = Font(size=9); F_DAY = Font(size=8)
FILL_HDR = PatternFill('solid', start_color='FFEDF2F9')
AL_C = Alignment(horizontal='center'); AL_L = Alignment(horizontal='left')

# 数据行读取（跳过空行/合计）
rows = []
for r in range(4, 71):
    wbs = fp.cell(row=r, column=1).value
    if not wbs: continue
    rows.append(dict(wbs=wbs, fpid=fp.cell(row=r, column=3).value, task=fp.cell(row=r, column=4).value,
                     ver=fp.cell(row=r, column=7).value, pd=fp.cell(row=r, column=8).value,
                     st=fp.cell(row=r, column=10).value, en=fp.cell(row=r, column=11).value))
chk(len(rows) == 67, '功能点行 %d≠67' % len(rows))

def build_sheet(name, idx, meta, data, note):
    if name in wb.sheetnames: wb.remove(wb[name])
    ws = wb.create_sheet(name, idx)
    n0 = len(meta)  # 元数据列数
    # 标题与口径
    ws.cell(1,1,name.split('-')[1] if False else '').value = None
    ws.cell(1,1).value = ('甘特图 · 任务级（%d 行全量）' if '任务' in name else '甘特图 · WBS × 版本（%d 行概览版）') % len(data)
    ws.cell(1,1).font = F_TITLE
    ws.cell(2,1).value = note; ws.cell(2,1).font = F_NOTE
    # 表头 + 月份带 + 日期行
    for c, h in enumerate(meta, 1):
        cell = ws.cell(3, c, h); cell.font = F_HDR; cell.fill = FILL_HDR; cell.alignment = AL_C
    prev = None
    for i, d in enumerate(AXIS):
        c = n0 + 1 + i
        cell = ws.cell(4, c, d.day); cell.font = F_DAY; cell.alignment = AL_C; cell.number_format = '0'
        cell.fill = FILL_HDR
        if prev is None or d.month != prev.month:
            mc = ws.cell(3, c, '%d月' % d.month); mc.font = F_HDR; mc.fill = FILL_HDR; mc.alignment = AL_C
        prev = d
    # 数据行（r5 起为前置段带）
    r = 5
    band = ws.cell(r, 1, '前置（09-02 ~ 09-18 · 需求+原型+底座）'); band.font = F_HDR; band.fill = FILL_HDR
    r += 1
    for row in data:
        vals = [row['wbs'], row.get('fpid',''), row['task'], row['ver'], row['pd'], row['st'], row['en']][:n0]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v); cell.font = F_DAT
            if isinstance(v, datetime.datetime): cell.number_format = 'MM-DD'
            if c in (1, 4, 5, 6, 7): cell.alignment = AL_C
        color = COLOR.get(row['ver'])
        if color and isinstance(row['st'], datetime.datetime):
            s = row['st'].date(); e = (row['en'] or row['st']).date() if isinstance(row['en'], datetime.datetime) else s
            bar = PatternFill('solid', start_color=color)
            for i, d in enumerate(AXIS):
                if s <= d <= e: ws.cell(r, n0+1+i).fill = bar
        r += 1
    # 列宽
    for c in range(1, n0+1):
        ws.column_dimensions[get_column_letter(c)].width = {1:5,2:9,3:44,4:7,5:6,6:8,7:8}.get(c, 9)
    for i in range(len(AXIS)):
        ws.column_dimensions[get_column_letter(n0+1+i)].width = 2.64
    ws.freeze_panes = ws.cell(5, n0+1)
    return ws

note = '口径：横轴 %d 个工作日（09-02~12-25·周末与国庆 10-01~07 不排）；09-15 增补口径（+12.5 人日·153.75），节点 10-20/11-24/12-14，人日待开发重估后刷新' % len(AXIS)
# 任务级（原索引 3）：A WBS/B 功能点/C 任务/D 版本/E 人日/F 开始/G 结束 + 轴
build_sheet('甘特图-任务级', 3, ['WBS','功能点','任务','版本','人日','开始','结束'], rows, note)
# 功能点级（原索引 2）：聚合
agg_pre = [dict(wbs='前置', fpid='', task='需求调研与 PRD（四轮走查+计费定案+评审冻结·09-18）', ver='前置', pd=6, st=D(9,2), en=D(9,18)),
           dict(wbs='前置', fpid='', task='原型改造与定稿（G31~G39·128+ HTML）', ver='前置', pd=1, st=D(9,14), en=D(9,18))]
b0 = [x for x in rows if x['wbs']=='B0']
agg_pre.append(dict(wbs='前置', fpid='', task='公共底座（脚手架/登录/通用组件）', ver='前置',
                    pd=sum(x['pd'] for x in b0 if x['ver']=='前置'),
                    st=min(x['st'] for x in b0 if x['ver']=='前置'), en=max(x['en'] for x in b0 if x['ver']=='前置')))
agg = [x for x in rows if x['wbs'] not in ('A1','A2','B0')]
groups = {}
for x in agg:
    k = (x['wbs'], x['ver'])
    g = groups.setdefault(k, dict(wbs=x['wbs'], fpid='', task=x['task'], ver=x['ver'], pd=0, st=None, en=None))
    g['pd'] += x['pd'] or 0
    if isinstance(x['st'], datetime.datetime):
        g['st'] = x['st'] if g['st'] is None else min(g['st'], x['st'])
    if isinstance(x['en'], datetime.datetime):
        g['en'] = x['en'] if g['en'] is None else max(g['en'], x['en'])
order = {w:i for i,w in enumerate(['B1','B2','B3','B4','B5','B6','B7','B8','C1','C2','C3'])}
vord = {'V1.0':0,'V1.1':1,'V2.0':2,'各版':3}
gl = [groups[k] for k in sorted(groups, key=lambda k:(order.get(k[0],99), vord.get(k[1],9)))]
ov_rows = agg_pre + gl
build_sheet('甘特图-功能点级', 2, ['WBS','模块 / 内容','人日','开始','结束'], ov_rows,
            note.replace('任务级','概览') if False else note)

wb.save(DST)
print('失败 %d 条' % len(fail))
for x in fail: print('  ' + x)
print('工作日轴：%d 天（%s ~ %s）' % (len(AXIS), AXIS[0], AXIS[-1]))
print('任务级行数 6+%d，功能点级 %d 行（前置 3 + 分组 %d）' % (len(rows), len(ov_rows), len(gl)))
print('已保存 ' + DST)
